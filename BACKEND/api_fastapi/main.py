import os
import datetime
import re
import unicodedata
from io import BytesIO
from decimal import Decimal
from hashlib import sha256
from pathlib import Path
from typing import List, Optional

import pandas as pd
import openpyxl
from fastapi import FastAPI, UploadFile, File, Form, Depends, HTTPException, status, Query
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from sqlalchemy import func
from sqlalchemy.orm import Session

from .database import get_db, engine
from .models_sqlalchemy import (
    Base,
    Analisis,
    CasoMorbilidad,
    CasoMortalidad,
    Paciente,
    Usuario,
    VMorbilidadCompleta,
    VMortalidadCompleta,
    CatTipoId,
    CatRemisiones,
    CatFuenteCausaMuerte,
    CatGrupoCausa,
    CatSitioDefuncion,
    CatConvivencia,
    CatEscolaridad,
    CatRegulacionFecundidad,
    CatMomentoMuerte,
    CatTipoParto,
    CatPersonalSalud,
    CatNivelAtencion,
)
from .schemas import (
    UsuarioRegister,
    UsuarioLogin,
    UsuarioResponse,
    AnalisisResponse,
    ClusteringRequest,
    PacienteResponse,
)
from .auth import verify_password, get_password_hash
from .sivigila_ingestion import persistir_dataframe_sivigila
from .processors import (
    MortalidadProcessor,
    MorbilidadProcessor,
    preparar_dataframe_analisis,
    es_valor_positivo,
)

# Initialize database schema (excluding views)
views_to_exclude = {'v_morbilidad_completa', 'v_mortalidad_completa'}
tables_to_create = [
    table for name, table in Base.metadata.tables.items()
    if name not in views_to_exclude
]
Base.metadata.create_all(bind=engine, tables=tables_to_create)

# Recreate database views dynamically based on connection dialect
from .database_views import create_database_views
db_session = Session(bind=engine)
try:
    create_database_views(db_session, engine.url.drivername)
finally:
    db_session.close()

app = FastAPI(
    title="MaternAnalytics API (FastAPI)",
    description="Backend de mortalidad y morbilidad materna migrado de Django a FastAPI.",
    version="1.0.0",
)

# CORS Middleware config (matches settings.py CORS_ALLOWED_ORIGINS)
app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173", "http://127.0.0.1:5173"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Ensure media directory exists
os.makedirs("media", exist_ok=True)
app.mount("/media", StaticFiles(directory="media"), name="media")

# --- Constants & Helper Logic (ported from api/views.py) ---
ERROR_ANALISIS_NO_ENCONTRADO = 'Análisis no encontrado.'
COLUMNA_MOMENTO_MUERTE = '9.1 Momento de la muerte'
COLUMNA_CAUSA_BASICA_CIE10 = '10.1 Causa básica CIE-10'
COLUMNA_TIPO_ID_MORTALIDAD = 'B. Tipo ID'
COLUMNA_NUMERO_ID_MORTALIDAD = 'C. Número ID'
COLUMNA_NUM_CPN_MORTALIDAD = '8.1 No. CPN'

COLUMNAS_MORTALIDAD = [
    'A. Nombres y Apellidos', COLUMNA_TIPO_ID_MORTALIDAD, COLUMNA_NUMERO_ID_MORTALIDAD,
    '5.1 Sitio de Defunción', '6.1 Convivencia', '6.3 Escolaridad',
    '6.4 Regulación Fecundidad', '6.5 Gestaciones', '6.6 Partos Vaginales',
    '6.7 Cesáreas', '6.8 Muertos', '6.9 Vivos', '6.10 Abortos',
    COLUMNA_NUM_CPN_MORTALIDAD, '8.2 Semana inicio CPN', COLUMNA_MOMENTO_MUERTE,
    '9.2 Semana gestación', '9.4 Tipo de parto', COLUMNA_CAUSA_BASICA_CIE10,
    '10.3.1 Demora 1', '10.3.2 Demora 2', '10.3.3 Demora 3', '10.3.4 Demora 4',
    'Fecha de Nacimiento',
]

COLUMNAS_MORBILIDAD = [
    'Nombres y apellidos', 'Tipo de ID', 'N° identificación',
    'N° gestaciones', 'Partos vaginales', 'Cesáreas', 'Abortos',
    'N° controles prenatales', 'Semanas inicio CPN',
    'Edad gestacional ocurrencia (sem)', 'Momento ocurrencia',
    'Eclampsia', 'Sepsis sistémica severa', 'Hemorragia obstétrica severa',
    'Preeclampsia', 'Ruptura uterina', 'Ingreso UCI', 'Cirugía adicional',
    'Transfusión', 'Total criterios', 'Causa principal CIE-10',
    'Días estancia hospitalaria', 'Días estancia UCI',
    'Fecha de Nacimiento', 'Fecha de egreso',
]

COLUMNAS_REQUERIDAS = {
    'mortalidad': COLUMNAS_MORTALIDAD,
    'morbilidad': COLUMNAS_MORBILIDAD,
}

ALIAS_COLUMNAS = {
    'morbilidad': {
        'Nombres y apellidos': ['Nombre y apellidos', 'Nombres y Apellidos'],
        'Tipo de ID': ['Tipo ID', 'Tipo identificación', 'Tipo de identificación'],
        'N° identificación': ['No identificación', 'Nro identificación', 'Nº identificación', 'Número identificación', 'Numero identificacion'],
        'N° gestaciones': ['No gestaciones', 'Nro gestaciones', 'Nº gestaciones', 'Numero gestaciones'],
        'Partos vaginales': ['Partos Vaginales'],
        'Cesáreas': ['Cesareas'],
        'N° controles prenatales': ['No controles prenatales', 'Nro controles prenatales', 'Nº controles prenatales', 'Numero controles prenatales'],
        'Causa principal CIE-10': ['Causa principal cie10', 'Causa principal CIE10'],
        'Días estancia hospitalaria': ['Dias estancia hospitalaria'],
        'Días estancia UCI': ['Dias estancia UCI'],
        'Fecha de Nacimiento': ['Fecha de nacimiento', 'Fecha nacimiento'],
        'Fecha de egreso': ['Fecha egreso', 'Fecha de egreso (dd/mm/aaaa)', 'Fecha egreso (dd/mm/aaaa)', 'Fecha de egreso (dd/mm/yyyy)', 'Fecha egreso (dd/mm/yyyy)'],
    },
    'mortalidad': {
        COLUMNA_TIPO_ID_MORTALIDAD: ['B. Tipo de ID', 'B Tipo ID'],
        COLUMNA_NUMERO_ID_MORTALIDAD: ['C. Numero ID', 'C Número ID'],
        COLUMNA_NUM_CPN_MORTALIDAD: ['8.1 N° CPN', '8.1 Nº CPN', '8.1 Numero CPN'],
        'Fecha de Nacimiento': ['Fecha de nacimiento', 'Fecha nacimiento', 'Fecha de Nacimiento (dd/mm/aaaa)', 'Fecha nacimiento (dd/mm/aaaa)'],
        '5.2 Fecha de defunción': [
            '5.2 Fecha defunción', 'Fecha de defunción', 'Fecha defunción',
            '5.2 Fecha de defunción (dd/mm/aaaa)', '5.2 Fecha de defuncion (dd/mm/aaaa)',
            '5.2 Fecha de defuncion', '5.2 Fecha defuncion',
            'Fecha de defuncion', 'Fecha defuncion',
        ],
    },
}

def obtener_limite_val(limit: Optional[int], default=100, maximo=500):
    if limit is None:
        return default
    return max(1, min(limit, maximo))

def serialize_row(row):
    d = {}
    for c in row.__table__.columns:
        val = getattr(row, c.name)
        if isinstance(val, (datetime.date, datetime.datetime, datetime.time)):
            d[c.name] = val.isoformat()
        elif isinstance(val, Decimal):
            d[c.name] = float(val) if val is not None else None
        else:
            d[c.name] = val
    return d

def _normalizar_encabezado(valor):
    texto = str(valor or '').strip().lower()
    texto = unicodedata.normalize('NFKD', texto).encode('ascii', 'ignore').decode('ascii')
    texto = texto.replace('n°', 'n ').replace('nº', 'n ').replace('no.', 'n ').replace('no ', 'n ')
    texto = re.sub(r'[^a-z0-9]+', ' ', texto)
    return ' '.join(texto.split())

def _construir_mapa_alias(tipo):
    mapa = {}
    for canonical in COLUMNAS_REQUERIDAS[tipo]:
        mapa[_normalizar_encabezado(canonical)] = canonical
    for canonical, alias_list in ALIAS_COLUMNAS.get(tipo, {}).items():
        mapa[_normalizar_encabezado(canonical)] = canonical
        for alias in alias_list:
            mapa[_normalizar_encabezado(alias)] = canonical
    return mapa

def _obtener_columnas_faltantes(tipo, columnas_archivo):
    mapa_alias = _construir_mapa_alias(tipo)
    presentes = {
        mapa_alias[normalizada]
        for columna in columnas_archivo
        for normalizada in [_normalizar_encabezado(columna)]
        if normalizada in mapa_alias
    }
    return [columna for columna in COLUMNAS_REQUERIDAS[tipo] if columna not in presentes]

def _canonizar_columnas_dataframe(df, tipo):
    mapa_alias = _construir_mapa_alias(tipo)
    renames = {}
    for columna in df.columns:
        canonical = mapa_alias.get(_normalizar_encabezado(columna))
        if canonical and canonical != columna and canonical not in df.columns:
            renames[columna] = canonical
    return df.rename(columns=renames)

def _seleccionar_fila_encabezados(filas, tipo):
    mejor_fila = []
    mejor_puntaje = -1
    for fila in filas:
        columnas = [str(valor).strip() for valor in fila if valor is not None and str(valor).strip()]
        if not columnas:
            continue
        puntaje = len(COLUMNAS_REQUERIDAS[tipo]) - len(_obtener_columnas_faltantes(tipo, columnas))
        if puntaje > mejor_puntaje:
            mejor_fila = columnas
            mejor_puntaje = puntaje
    return mejor_fila

def _seleccionar_hoja_y_encabezados_openpyxl(workbook, tipo):
    mejor_fila = []
    mejor_puntaje = -1
    mejor_hoja = workbook.active.title
    for hoja in workbook.worksheets:
        filas = [
            [cell.value for cell in row]
            for row in hoja.iter_rows(min_row=1, max_row=5)
        ]
        fila = _seleccionar_fila_encabezados(filas, tipo)
        puntaje = len(COLUMNAS_REQUERIDAS[tipo]) - len(_obtener_columnas_faltantes(tipo, fila)) if fila else -1
        if puntaje > mejor_puntaje:
            mejor_fila = fila
            mejor_puntaje = puntaje
            mejor_hoja = hoja.title
    return mejor_hoja, mejor_fila

def leer_columnas_excel(file_like, tipo):
    try:
        wb = openpyxl.load_workbook(file_like, read_only=True, data_only=True)
        _, headers = _seleccionar_hoja_y_encabezados_openpyxl(wb, tipo)
        wb.close()
        return headers
    except Exception:
        return None

def _detectar_indice_encabezados_dataframe(df_sin_encabezado, tipo):
    mejor_indice = 0
    mejor_puntaje = -1
    for indice in range(min(5, len(df_sin_encabezado.index))):
        fila = df_sin_encabezado.iloc[indice].tolist()
        columnas = [str(valor).strip() for valor in fila if not pd.isna(valor) and str(valor).strip()]
        if not columnas:
            continue
        puntaje = len(COLUMNAS_REQUERIDAS[tipo]) - len(_obtener_columnas_faltantes(tipo, columnas))
        if puntaje > mejor_puntaje:
            mejor_indice = indice
            mejor_puntaje = puntaje
    return mejor_indice

def _seleccionar_hoja_y_encabezados_dataframe(sheets, tipo):
    mejor_hoja = None
    mejor_indice = 0
    mejor_puntaje = -1
    for nombre_hoja, dataframe in sheets.items():
        indice = _detectar_indice_encabezados_dataframe(dataframe, tipo)
        fila = dataframe.iloc[indice].tolist() if len(dataframe.index) > indice else []
        columnas = [str(valor).strip() for valor in fila if not pd.isna(valor) and str(valor).strip()]
        puntaje = len(COLUMNAS_REQUERIDAS[tipo]) - len(_obtener_columnas_faltantes(tipo, columnas)) if columnas else -1
        if puntaje > mejor_puntaje:
            mejor_hoja = nombre_hoja
            mejor_indice = indice
            mejor_puntaje = puntaje
    return mejor_hoja, mejor_indice

def _leer_dataframe_excel(file_like, tipo):
    file_like.seek(0)
    hojas = pd.read_excel(file_like, engine='openpyxl', header=None, sheet_name=None)
    nombre_hoja, indice_encabezados = _seleccionar_hoja_y_encabezados_dataframe(hojas, tipo)
    file_like.seek(0)
    df = pd.read_excel(file_like, engine='openpyxl', header=indice_encabezados, sheet_name=nombre_hoja)
    return _canonizar_columnas_dataframe(df, tipo)

def calcular_resumen_desde_dataframe(df, tipo):
    try:
        df, limpieza = preparar_dataframe_analisis(df)
        resumen = {
            'total_registros': len(df),
            'total_registros_original': limpieza['total_original'],
            'filas_vacias_omitidas': limpieza['filas_vacias_omitidas'],
            'filas_duplicadas_omitidas': limpieza['filas_duplicadas_omitidas'],
        }
        if tipo == 'mortalidad':
            if COLUMNA_MOMENTO_MUERTE in df.columns:
                resumen['distribucion_momento'] = (
                    df[COLUMNA_MOMENTO_MUERTE].value_counts().to_dict()
                )
            if COLUMNA_CAUSA_BASICA_CIE10 in df.columns:
                top5 = df[COLUMNA_CAUSA_BASICA_CIE10].value_counts().head(5).to_dict()
                resumen['top5_causas'] = {str(k): int(v) for k, v in top5.items()}
        elif tipo == 'morbilidad':
            criterios_cols = ['Eclampsia', 'Sepsis sistémica severa',
                              'Hemorragia obstétrica severa', 'Preeclampsia', 'Ruptura uterina']
            presentes = [c for c in criterios_cols if c in df.columns]
            if presentes:
                resumen['criterios_frecuencia'] = {
                    c: int(df[c].apply(es_valor_positivo).sum()) for c in presentes
                }
        return resumen
    except Exception:
        return {}

MORTALIDAD_ANALISIS_MAPPING = {
    'nombres_apellidos': 'A. Nombres y Apellidos',
    'tipo_id': 'B. Tipo ID',
    'numero_id': 'C. Número ID',
    'sitio_defuncion': '5.1 Sitio de Defunción',
    'fecha_defuncion': '5.2 Fecha de defunción',
    'fecha_parto': '9.3 Fecha parto (dd/mm/aaaa)',
    'fecha_nacimiento': 'Fecha de Nacimiento',
    'edad': 'Edad',
    'convivencia': '6.1 Convivencia',
    'escolaridad': '6.3 Escolaridad',
    'regulacion_fecundidad': '6.4 Regulación Fecundidad',
    'gestaciones': '6.5 Gestaciones',
    'partos_vaginales': '6.6 Partos Vaginales',
    'cesareas': '6.7 Cesáreas',
    'nacidos_muertos': '6.8 Muertos',
    'hijos_vivos': '6.9 Vivos',
    'abortos': '6.10 Abortos',
    'num_cpn': '8.1 No. CPN',
    'semana_inicio_cpn': '8.2 Semana inicio CPN',
    'momento_muerte': '9.1 Momento de la muerte',
    'semana_gestacion_muerte': '9.2 Semana gestación',
    'tipo_parto': '9.4 Tipo de parto',
    'causa_basica_cie10': '10.1 Causa básica CIE-10',
    'demora_1': '10.3.1 Demora 1',
    'demora_2': '10.3.2 Demora 2',
    'demora_3': '10.3.3 Demora 3',
    'demora_4': '10.3.4 Demora 4',
}

def construir_dataframe_analisis_desde_bd(db, tipo):
    if tipo != 'mortalidad':
        return None
    registros = db.query(VMortalidadCompleta).order_by(VMortalidadCompleta.id_caso).all()
    if not registros:
        return pd.DataFrame(columns=MORTALIDAD_ANALISIS_MAPPING.values())
    records = []
    for r in registros:
        d = {}
        for key in MORTALIDAD_ANALISIS_MAPPING.keys():
            val = getattr(r, key)
            if isinstance(val, Decimal):
                d[key] = float(val) if val is not None else None
            elif isinstance(val, (datetime.date, datetime.datetime)):
                d[key] = val.isoformat()
            else:
                d[key] = val
        records.append(d)
    dataframe = pd.DataFrame.from_records(records)
    return dataframe.rename(columns=MORTALIDAD_ANALISIS_MAPPING)

def construir_archivo_analisis_desde_dataframe(df_fuente, tipo, nombre_archivo):
    df_analisis, _ = preparar_dataframe_analisis(df_fuente)
    resumen = calcular_resumen_desde_dataframe(df_analisis, tipo)
    total_registros = resumen.pop('total_registros', len(df_analisis))

    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
        df_analisis.to_excel(writer, index=False)

    contenido = buffer.getvalue()
    
    # Save physical file
    now = datetime.datetime.now()
    upload_dir = Path("media") / "uploads" / f"{now.year}" / f"{now.month:02d}"
    upload_dir.mkdir(parents=True, exist_ok=True)
    
    safe_name = "".join([c if c.isalnum() or c in (".", "_", "-") else "_" for c in nombre_archivo])
    file_path = upload_dir / safe_name
    with open(file_path, "wb") as f:
        f.write(contenido)

    db_relative_path = f"/media/uploads/{now.year}/{now.month:02d}/{safe_name}"
    archivo_hash = sha256(contenido).hexdigest()

    return db_relative_path, archivo_hash, resumen, total_registros

def construir_archivo_analisis_acumulado(db, analisis_existente, df_nuevo, tipo, nombre_archivo):
    dataframes = []
    if analisis_existente is not None and analisis_existente.archivo:
        try:
            # Convert DB relative path (e.g. /media/...) to filesystem path
            # Remove leading slash if present
            clean_path = analisis_existente.archivo.lstrip("/")
            full_path = Path(clean_path)
            if full_path.exists():
                df_anterior = pd.read_excel(full_path, engine='openpyxl')
                dataframes.append(df_anterior)
        except Exception as e:
            print(f"Error reading existing file: {e}")

    dataframes.append(df_nuevo.copy())
    df_acumulado = pd.concat(dataframes, ignore_index=True) if len(dataframes) > 1 else dataframes[0]
    return construir_archivo_analisis_desde_dataframe(df_acumulado, tipo, nombre_archivo)

def obtener_analisis_unicos(lista):
    unicos = {}
    for analisis in lista:
        clave = analisis.tipo
        existente = unicos.get(clave)
        if existente is None or analisis.fecha_carga > existente.fecha_carga:
            unicos[clave] = analisis
    return sorted(unicos.values(), key=lambda a: a.fecha_carga, reverse=True)

def _parse_valor_fecha(v):
    """Parsea un valor individual (mantenido por compatibilidad)."""
    if v is None:
        return pd.NaT
    if hasattr(v, 'year'):
        return pd.Timestamp(v)
    try:
        n = float(v)
        if 1000 < n < 100000:
            return pd.Timestamp('1899-12-30') + pd.Timedelta(days=int(n))
    except (ValueError, TypeError):
        pass
    try:
        return pd.to_datetime(str(v), dayfirst=True, errors='coerce')
    except Exception:
        return pd.NaT

def _parse_serie_fechas(serie):
    """Parsea una serie de fechas de forma vectorizada y eficiente (datetime, serial Excel o string DD/MM/YYYY)."""
    if pd.api.types.is_datetime64_any_dtype(serie):
        return serie

    # 1. Identificar y convertir seriales de Excel (rango 1000 - 100000)
    s_numeric = pd.to_numeric(serie, errors='coerce')
    is_excel_serial = (s_numeric > 1000) & (s_numeric < 100000)
    
    serie_clean = serie.copy()
    if is_excel_serial.any():
        if isinstance(serie_clean, pd.Series):
            serie_clean = serie_clean.mask(is_excel_serial)
        else:
            serie_clean = pd.Series(serie_clean).mask(is_excel_serial)

    # 2. Intentar convertir directamente con dayfirst=True para evitar warnings con DD/MM/YYYY
    import warnings
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", category=UserWarning, message=".*Parsing dates.*")
        fechas = pd.to_datetime(serie_clean, dayfirst=True, errors='coerce')
    
    # 3. Combinar con seriales de Excel
    if is_excel_serial.any():
        excel_days = s_numeric[is_excel_serial].astype(int)
        fechas_excel = pd.to_datetime(excel_days, unit='D', origin='1899-12-30', errors='coerce')
        fechas = fechas.fillna(fechas_excel)
        
    # 4. Para cualquier valor restante que falló y es no nulo, intentar sin dayfirst
    por_vias_alternas = fechas.isna() & serie_clean.notna()
    if por_vias_alternas.any():
        alternas_str = serie_clean[por_vias_alternas].astype(str)
        fechas_alt = pd.to_datetime(alternas_str, errors='coerce')
        fechas = fechas.fillna(fechas_alt)

    return fechas

def _candidatos_fecha(tipo):
    if tipo == 'mortalidad':
        return ['5.2 Fecha de defunción', '5.2 Fecha de defuncion', 'Fecha de defunción', 'Fecha de defuncion']
    else:
        return ['Fecha de egreso', 'Fecha egreso']

def _detectar_col_fecha(df, tipo):
    candidatos = _candidatos_fecha(tipo)
    for col in df.columns:
        if col in candidatos:
            return col
    for col in df.columns:
        if 'fecha' in str(col).lower():
            return col
    return None

def _extraer_anos_disponibles(df, tipo):
    col = _detectar_col_fecha(df, tipo)
    if col is None:
        return []
    fechas = _parse_serie_fechas(df[col])
    anos = fechas.dt.year.dropna().unique().astype(int).tolist()
    return sorted(anos, reverse=True)

def _enriquecer_df_con_fecha(db: Session, df, tipo):
    col = _detectar_col_fecha(df, tipo)
    if col is not None and df[col].notna().any():
        return df

    # Enrich from database records matching by type and number_id
    if tipo == 'mortalidad':
        casos = db.query(CasoMortalidad.id_caso, Paciente.numero_id, CasoMortalidad.fecha_defuncion)\
                  .join(Paciente, CasoMortalidad.id_paciente == Paciente.id_paciente).all()
        fechas_map = {c.numero_id: c.fecha_defuncion for c in casos if c.numero_id}
        col_id = 'C. Número ID'
        col_fecha_new = '5.2 Fecha de defunción'
    else:
        casos = db.query(CasoMorbilidad.id_caso, Paciente.numero_id, CasoMorbilidad.fecha_egreso)\
                  .join(Paciente, CasoMorbilidad.id_paciente == Paciente.id_paciente).all()
        fechas_map = {c.numero_id: c.fecha_egreso for c in casos if c.numero_id}
        col_id = 'N° identificación'
        col_fecha_new = 'Fecha de egreso'

    if col_id in df.columns:
        df[col_fecha_new] = df[col_id].astype(str).map(fechas_map)
        df[col_fecha_new] = pd.to_datetime(df[col_fecha_new], errors='coerce')
    
    return df

def _filtrar_dataframe_por_fecha(df, tipo, year, month):
    col = _detectar_col_fecha(df, tipo)
    if col is None:
        return df
    fechas = _parse_serie_fechas(df[col])
    mask = pd.Series([True] * len(df), index=df.index)
    if year:
        mask &= fechas.dt.year == int(year)
    if month:
        mask &= fechas.dt.month == int(month)
    return df[mask].reset_index(drop=True)


MESES_ABREV = ['Ene', 'Feb', 'Mar', 'Abr', 'May', 'Jun',
               'Jul', 'Ago', 'Sep', 'Oct', 'Nov', 'Dic']

def _sanitize_json(obj):
    """Reemplaza nan/inf de numpy/pandas con None para que json.dumps no falle."""
    import math
    if isinstance(obj, dict):
        return {k: _sanitize_json(v) for k, v in obj.items()}
    if isinstance(obj, list):
        return [_sanitize_json(v) for v in obj]
    if isinstance(obj, float) and (math.isnan(obj) or math.isinf(obj)):
        return None
    return obj

def _calcular_distribucion_mensual(df, tipo):
    col = _detectar_col_fecha(df, tipo)
    if col is None:
        return {}
    fechas = _parse_serie_fechas(df[col]).dropna()
    if fechas.empty:
        return {}
    conteos = fechas.dt.month.value_counts().sort_index()
    return {
        'labels': [MESES_ABREV[m - 1] for m in conteos.index],
        'values': [int(v) for v in conteos.values],
    }


# --- REST API Endpoints ---

# Auth Endpoints
@app.post("/api/auth/register/", response_model=UsuarioResponse, status_code=status.HTTP_201_CREATED)
def register_usuario(user_in: UsuarioRegister, db: Session = Depends(get_db)):
    user_in.email = user_in.email.strip().lower()
    if len(user_in.password) < 6:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="La contraseña debe tener al menos 6 caracteres."
        )
    
    existing = db.query(Usuario).filter(Usuario.email == user_in.email).first()
    if existing:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Ya existe una cuenta con este correo electrónico."
        )

    nuevo_usuario = Usuario(
        nombre=user_in.nombre.strip(),
        email=user_in.email,
        password_hash=get_password_hash(user_in.password),
        fecha_registro=datetime.datetime.utcnow()
    )
    db.add(nuevo_usuario)
    db.commit()
    db.refresh(nuevo_usuario)
    return nuevo_usuario

@app.post("/api/auth/login/", response_model=UsuarioResponse)
def login_usuario(credentials: UsuarioLogin, db: Session = Depends(get_db)):
    credentials.email = credentials.email.strip().lower()
    user = db.query(Usuario).filter(Usuario.email == credentials.email).first()
    if not user:
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Correo o contraseña incorrectos."
        )
    
    if not verify_password(credentials.password, user.password_hash):
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Correo o contraseña incorrectos."
        )
    
    return user


# Ingestion & Uploads
def handle_file_upload(tipo: str, archivo: UploadFile, db: Session):
    # Validate headers
    archivo.file.seek(0)
    columnas_archivo = leer_columnas_excel(archivo.file, tipo)
    if columnas_archivo is None:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="No se pudo leer el archivo. Verifica que sea un Excel válido."
        )

    faltantes = _obtener_columnas_faltantes(tipo, columnas_archivo)
    if faltantes:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail={"error": "Faltan columnas requeridas.", "columnas_faltantes": faltantes}
        )

    try:
        df = _leer_dataframe_excel(archivo.file, tipo)
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"No se pudo leer el contenido del Excel: {str(e)}"
        )

    df_cleaned, _ = preparar_dataframe_analisis(df)

    try:
        # DB Transactions in SQLAlchemy
        db.begin_nested() # Create a savepoint
        
        analisis_existente = db.query(Analisis).filter(Analisis.tipo == tipo)\
                               .order_by(Analisis.fecha_carga.desc(), Analisis.id.desc()).first()
        
        persistencia = persistir_dataframe_sivigila(db, df_cleaned, tipo)

        df_autoritativo = construir_dataframe_analisis_desde_bd(db, tipo)
        if df_autoritativo is not None:
            archivo_analisis, archivo_hash, resumen, total_registros = construir_archivo_analisis_desde_dataframe(
                df_autoritativo,
                tipo,
                archivo.filename,
            )
        else:
            archivo_analisis, archivo_hash, resumen, total_registros = construir_archivo_analisis_acumulado(
                db,
                analisis_existente,
                df_cleaned,
                tipo,
                archivo.filename,
            )

        if analisis_existente is None:
            analisis = Analisis(
                tipo=tipo,
                nombre_archivo=archivo.filename,
                archivo_hash=archivo_hash,
                archivo=archivo_analisis,
                total_registros=total_registros,
                resumen=resumen,
                fecha_carga=datetime.datetime.utcnow()
            )
            db.add(analisis)
        else:
            analisis_existente.nombre_archivo = archivo.filename
            analisis_existente.archivo_hash = archivo_hash
            analisis_existente.archivo = archivo_analisis
            analisis_existente.total_registros = total_registros
            analisis_existente.resumen = resumen
            analisis_existente.fecha_carga = datetime.datetime.utcnow()
            analisis = analisis_existente
            
        db.commit()
    except ValueError as e:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=str(e)
        )
    except Exception as e:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"No se pudo persistir el archivo en SIVIGILA: {str(e)}"
        )

    # Return structure matching django response
    res = {
        "id": analisis.id,
        "tipo": analisis.tipo,
        "nombre_archivo": analisis.nombre_archivo,
        "archivo": analisis.archivo,
        "fecha_carga": analisis.fecha_carga.isoformat(),
        "total_registros": analisis.total_registros,
        "resumen": analisis.resumen,
        "sivigila": persistencia
    }
    return res

@app.post("/api/subir/")
def subir_archivo(tipo: str = Form(...), archivo: UploadFile = File(...), db: Session = Depends(get_db)):
    return handle_file_upload(tipo, archivo, db)

@app.post("/api/analisis/")
def crear_analisis(tipo: str = Form(...), archivo: UploadFile = File(...), db: Session = Depends(get_db)):
    return handle_file_upload(tipo, archivo, db)


# Analyses query endpoints
@app.get("/api/analisis/", response_model=List[AnalisisResponse])
def listar_analisis(db: Session = Depends(get_db)):
    analisis_list = db.query(Analisis).order_by(Analisis.fecha_carga.desc()).all()
    # Filter to only unique types of latest date
    unique_list = obtener_analisis_unicos(analisis_list)
    return unique_list

@app.get("/api/analisis/{pk}/", response_model=AnalisisResponse)
def detalle_analisis(pk: int, db: Session = Depends(get_db)):
    analisis = db.query(Analisis).filter(Analisis.id == pk).first()
    if not analisis:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=ERROR_ANALISIS_NO_ENCONTRADO
        )
    return analisis


# Complex Analytics
@app.get("/api/analisis/{pk}/completo/")
def analisis_completo(
    pk: int,
    year: Optional[str] = Query(None),
    month: Optional[str] = Query(None),
    db: Session = Depends(get_db)
):
    analisis = db.query(Analisis).filter(Analisis.id == pk).first()
    if not analisis:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=ERROR_ANALISIS_NO_ENCONTRADO
        )

    # Clean path (strip leading slash)
    clean_path = analisis.archivo.lstrip("/")
    full_path = Path(clean_path)
    if not full_path.exists():
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="El archivo físico del análisis no existe en el servidor."
        )

    try:
        df = pd.read_excel(full_path, engine='openpyxl')
        df = _canonizar_columnas_dataframe(df, analisis.tipo)
        df, limpieza = preparar_dataframe_analisis(df)
        df = _enriquecer_df_con_fecha(db, df, analisis.tipo)
        anos_disponibles = _extraer_anos_disponibles(df, analisis.tipo)
        df = _filtrar_dataframe_por_fecha(df, analisis.tipo, year, month)

        if analisis.tipo == 'mortalidad':
            processor = MortalidadProcessor(df)
            resultado = {
                'tipo': 'mortalidad',
                'id': analisis.id,
                'nombre_archivo': analisis.nombre_archivo,
                'fecha_carga': analisis.fecha_carga.isoformat(),
                'limpieza_datos': limpieza,
                'estadisticas_basicas': processor.calcular_estadisticas_basicas(),
                'momento_muerte': processor.analizar_momento_muerte(),
                'demoras': processor.analizar_demoras(),
                'causas_cie10': processor.analizar_causas_cie10(top_n=15),
                'obstetrico_edad': processor.analizar_obstetrico_por_edad(),
                'anos_disponibles': anos_disponibles,
                'distribucion_mensual': _calcular_distribucion_mensual(df, 'mortalidad'),
                'filtros_activos': {'year': year, 'month': month},
            }
        else:  # morbilidad
            processor = MorbilidadProcessor(df)
            resultado = {
                'tipo': 'morbilidad',
                'id': analisis.id,
                'nombre_archivo': analisis.nombre_archivo,
                'fecha_carga': analisis.fecha_carga.isoformat(),
                'limpieza_datos': limpieza,
                'estadisticas_basicas': processor.calcular_estadisticas_basicas(),
                'criterios_inclusion': processor.analizar_criterios_inclusion(),
                'momento_ocurrencia': processor.analizar_momento_ocurrencia(),
                'institucion_referencia': processor.analizar_institucion_referencia(),
                'tiempo_remision': processor.analizar_tiempo_remision(),
                'obstetrico_edad': processor.analizar_obstetrico_por_edad(),
                'anos_disponibles': anos_disponibles,
                'distribucion_mensual': _calcular_distribucion_mensual(df, 'morbilidad'),
                'filtros_activos': {'year': year, 'month': month},
            }
        return _sanitize_json(resultado)
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Error al procesar el análisis: {str(e)}"
        )

@app.post("/api/analisis/{pk}/clustering/")
def clustering_analisis(pk: int, req: ClusteringRequest, db: Session = Depends(get_db)):
    analisis = db.query(Analisis).filter(Analisis.id == pk).first()
    if not analisis:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=ERROR_ANALISIS_NO_ENCONTRADO
        )
    
    clean_path = analisis.archivo.lstrip("/")
    full_path = Path(clean_path)
    if not full_path.exists():
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="El archivo físico del análisis no existe en el servidor."
        )

    try:
        df = pd.read_excel(full_path, engine='openpyxl')
        df, limpieza = preparar_dataframe_analisis(df)
        
        if analisis.tipo == 'mortalidad':
            processor = MortalidadProcessor(df)
            if req.tipo_clustering == 'jerarquico':
                resultado = processor.clustering_jerarquico()
            else:
                resultado = processor.clustering_factores_riesgo(n_clusters=req.n_clusters)
        else:
            processor = MorbilidadProcessor(df)
            resultado = processor.clustering_perfiles_morbilidad(n_clusters=req.n_clusters)
            
        resultado['analisis_id'] = analisis.id
        resultado['tipo_analisis'] = analisis.tipo
        resultado['tipo_clustering'] = req.tipo_clustering
        resultado['limpieza_datos'] = limpieza
        
        return _sanitize_json(resultado)
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Error en clustering: {str(e)}"
        )

@app.get("/api/analisis/{pk}/heatmap/")
def heatmap_correlacion(pk: int, db: Session = Depends(get_db)):
    analisis = db.query(Analisis).filter(Analisis.id == pk).first()
    if not analisis:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=ERROR_ANALISIS_NO_ENCONTRADO
        )
    
    if analisis.tipo != 'morbilidad':
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Heatmap solo disponible para análisis de morbilidad"
        )
        
    clean_path = analisis.archivo.lstrip("/")
    full_path = Path(clean_path)
    if not full_path.exists():
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="El archivo físico del análisis no existe en el servidor."
        )

    try:
        df = pd.read_excel(full_path, engine='openpyxl')
        df, limpieza = preparar_dataframe_analisis(df)
        processor = MorbilidadProcessor(df)
        resultado = processor.heatmap_correlacion()
        
        resultado['analisis_id'] = analisis.id
        resultado['limpieza_datos'] = limpieza

        return _sanitize_json(resultado)
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Error al generar heatmap: {str(e)}"
        )


@app.get("/api/analisis/{pk}/extra-columna/")
def extra_columna_analisis(pk: int, db: Session = Depends(get_db)):
    analisis = db.query(Analisis).filter(Analisis.id == pk).first()
    if not analisis:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=ERROR_ANALISIS_NO_ENCONTRADO
        )
    if analisis.tipo != 'mortalidad':
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Extra columna solo disponible para análisis de mortalidad"
        )

    clean_path = analisis.archivo.lstrip("/")
    full_path = Path(clean_path)
    if not full_path.exists():
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="El archivo físico del análisis no existe en el servidor."
        )

    try:
        df = pd.read_excel(full_path, engine='openpyxl')
        df = _canonizar_columnas_dataframe(df, 'mortalidad')
        df, _ = preparar_dataframe_analisis(df)

        COLUMNAS_EXTRA = [
            ('5.1 Sitio de Defunción', 'Distribución por Sitio de Defunción', 'Sitio', 'Casos'),
            ('6.3 Escolaridad', 'Distribución por Nivel de Escolaridad', 'Escolaridad', 'Casos'),
            ('6.1 Convivencia', 'Distribución por Convivencia', 'Convivencia', 'Casos'),
        ]

        for col, titulo, x_title, y_title in COLUMNAS_EXTRA:
            if col not in df.columns:
                continue
            serie = df[col].dropna().astype(str).str.strip()
            serie = serie[~serie.str.lower().isin({'', 'nan', 'none', 'null'})]
            if len(serie) == 0:
                continue
            conteos = serie.value_counts()
            return {
                'chart': {
                    'title': titulo,
                    'subtitle': f'Total: {len(serie)} casos con dato registrado',
                    'labels': conteos.index.tolist(),
                    'values': [int(v) for v in conteos.values],
                    'total': int(len(serie)),
                    'xTitle': x_title,
                    'yTitle': y_title,
                    'orientation': 'h',
                }
            }

        return {'chart': None}
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Error al procesar extra columna: {str(e)}"
        )


# Sivigila DB Querying Endpoints
@app.get("/api/sivigila/resumen/")
def resumen_sivigila(db: Session = Depends(get_db)):
    return {
        'pacientes': db.query(Paciente).count(),
        'casos_morbilidad': db.query(CasoMorbilidad).count(),
        'casos_mortalidad': db.query(CasoMortalidad).count(),
    }

@app.get("/api/sivigila/pacientes/")
def listar_pacientes(limit: Optional[int] = Query(None), db: Session = Depends(get_db)):
    limite = obtener_limite_val(limit)
    results = db.query(Paciente, CatTipoId)\
                .outerjoin(CatTipoId, Paciente.id_tipo_id == CatTipoId.id)\
                .order_by(Paciente.id_paciente)\
                .limit(limite).all()
    
    out = []
    for p, t in results:
        out.append({
            'id_paciente': p.id_paciente,
            'nombres_apellidos': p.nombres_apellidos,
            'tipo_identificacion': t.codigo if t else None,
            'descripcion_tipo_identificacion': t.descripcion if t else None,
            'numero_id': p.numero_id,
            'fecha_nacimiento': p.fecha_nacimiento.isoformat() if p.fecha_nacimiento else None,
            'creado_en': p.creado_en.isoformat() if p.creado_en else None
        })
    return out

@app.get("/api/sivigila/morbilidad/")
def listar_morbilidad_sivigila(limit: Optional[int] = Query(None), db: Session = Depends(get_db)):
    limite = obtener_limite_val(limit)
    try:
        casos = db.query(VMorbilidadCompleta).order_by(VMorbilidadCompleta.id_caso).limit(limite).all()
        return [serialize_row(caso) for caso in casos]
    except Exception as e:
        # If the view isn't created in PG yet, return a clean warning so front doesn't crash completely.
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Error al consultar la vista v_morbilidad_completa. Asegúrate de haberla creado en la base de datos: {str(e)}"
        )

@app.get("/api/sivigila/mortalidad/")
def listar_mortalidad_sivigila(limit: Optional[int] = Query(None), db: Session = Depends(get_db)):
    limite = obtener_limite_val(limit)
    try:
        casos = db.query(VMortalidadCompleta).order_by(VMortalidadCompleta.id_caso).limit(limite).all()
        return [serialize_row(caso) for caso in casos]
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Error al consultar la vista v_mortalidad_completa. Asegúrate de haberla creado en la base de datos: {str(e)}"
        )
